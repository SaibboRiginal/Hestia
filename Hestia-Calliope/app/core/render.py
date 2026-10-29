"""Blocks → PDF (reportlab), Word (python-docx), Excel (openpyxl) with three shared themes, and the final check:
the file opens again, has pages/sheets, and contains the title and every «must contain» text."""
from __future__ import annotations

import io
import re
from datetime import date
from xml.sax.saxutils import escape

from .parse import plain, text_of

THEMES = {
    "classico": {"label": "Classico", "font": "Times", "accent": "1F3A5F", "muted": "5B6573", "fill": "E8EDF3"},
    "moderno": {"label": "Moderno", "font": "Helvetica", "accent": "0F766E", "muted": "52606D", "fill": "E6F4F1"},
    "sobrio": {"label": "Sobrio", "font": "Helvetica", "accent": "2B2B2B", "muted": "6B6B6B", "fill": "EFEFEF"},
}
FORMATS = {"pdf": ("application/pdf", "pdf"),
           "docx": ("application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx"),
           "xlsx": ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx")}


class RenderError(RuntimeError):
    pass


def _pdf_fonts(font: str) -> tuple[str, str, str, str]:
    if font == "Times":
        return "Times-Roman", "Times-Bold", "Times-Italic", "Times-BoldItalic"
    return "Helvetica", "Helvetica-Bold", "Helvetica-Oblique", "Helvetica-BoldOblique"


def _markup(runs) -> str:
    out = []
    for text, b, i in runs:
        t = "<br/>" if text == "\n" else escape(text)
        if i:
            t = f"<i>{t}</i>"
        if b:
            t = f"<b>{t}</b>"
        out.append(t)
    return "".join(out)


def to_pdf(blocks: list[dict], title: str, theme: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    reg, bold, ital, _ = _pdf_fonts(theme["font"])
    accent, muted, fill = (colors.HexColor("#" + theme[k]) for k in ("accent", "muted", "fill"))
    body = ParagraphStyle("body", fontName=reg, fontSize=10.5, leading=15, spaceAfter=6)
    heads = {1: ParagraphStyle("h1", fontName=bold, fontSize=16, leading=20, textColor=accent, spaceBefore=10,
                               spaceAfter=6),
             2: ParagraphStyle("h2", fontName=bold, fontSize=13, leading=17, textColor=accent, spaceBefore=8,
                               spaceAfter=4),
             3: ParagraphStyle("h3", fontName=bold, fontSize=11.5, leading=15, textColor=muted, spaceBefore=6,
                               spaceAfter=3)}
    cover = ParagraphStyle("title", fontName=bold, fontSize=22, leading=26, textColor=accent, spaceAfter=4)
    sub = ParagraphStyle("sub", fontName=ital, fontSize=9, textColor=muted, spaceAfter=14)
    story = [Paragraph(escape(title), cover), Paragraph(date.today().strftime("%d/%m/%Y"), sub)]
    for b in blocks:
        if b["type"] == "h":
            story.append(Paragraph(_markup(b["runs"]), heads[b["level"]]))
        elif b["type"] == "p":
            story.append(Paragraph(_markup(b["runs"]), body))
        elif b["type"] == "list":
            story.append(ListFlowable([ListItem(Paragraph(_markup(i), body), leftIndent=12) for i in b["items"]],
                                      bulletType="1" if b["ordered"] else "bullet", leftIndent=14,
                                      bulletFontName=reg, bulletColor=accent))
        elif b["type"] == "table":
            width = max(len(r) for r in b["rows"])
            data = [[Paragraph(_markup(c["runs"]), body) for c in r] + [""] * (width - len(r)) for r in b["rows"]]
            t = Table(data, repeatRows=1 if any(c["header"] for c in b["rows"][0]) else 0, hAlign="LEFT")
            style = [("GRID", (0, 0), (-1, -1), 0.4, muted), ("VALIGN", (0, 0), (-1, -1), "TOP")]
            if any(c["header"] for c in b["rows"][0]):
                style.append(("BACKGROUND", (0, 0), (-1, 0), fill))
            t.setStyle(TableStyle(style))
            story += [t, Spacer(1, 8)]
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, title=title, leftMargin=2.2 * cm, rightMargin=2.2 * cm,
                            topMargin=2 * cm, bottomMargin=2 * cm)

    def _footer(canvas, d):
        canvas.saveState()
        canvas.setFont(reg, 8)
        canvas.setFillColor(muted)
        canvas.drawRightString(A4[0] - 2.2 * cm, 1.2 * cm, str(d.page))
        canvas.restoreState()

    doc.build(story, onFirstPage=_footer, onLaterPages=_footer)
    return buf.getvalue()


def to_docx(blocks: list[dict], title: str, theme: dict) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    accent = RGBColor.from_string(theme["accent"])
    font = "Times New Roman" if theme["font"] == "Times" else "Arial"
    doc.styles["Normal"].font.name = font
    doc.styles["Normal"].font.size = Pt(11)
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        doc.styles[name].font.name = font
        doc.styles[name].font.color.rgb = accent
    doc.add_heading(title, level=0)

    def runs_into(par, runs):
        for text, b, i in runs:
            if text == "\n":
                par.add_run().add_break()
                continue
            r = par.add_run(text)
            r.bold, r.italic = b or None, i or None

    for b in blocks:
        if b["type"] == "h":
            runs_into(doc.add_heading("", level=b["level"]), b["runs"])
        elif b["type"] == "p":
            runs_into(doc.add_paragraph(), b["runs"])
        elif b["type"] == "list":
            for it in b["items"]:
                runs_into(doc.add_paragraph(style="List Number" if b["ordered"] else "List Bullet"), it)
        elif b["type"] == "table":
            width = max(len(r) for r in b["rows"])
            table = doc.add_table(rows=len(b["rows"]), cols=width)
            table.style = "Table Grid"
            for ri, row in enumerate(b["rows"]):
                for ci, c in enumerate(row):
                    par = table.cell(ri, ci).paragraphs[0]
                    runs_into(par, [(t, bb or c["header"], ii) for t, bb, ii in c["runs"]])
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _cell_value(text: str):
    t = text.strip()
    if re.fullmatch(r"-?\d+", t):
        return int(t)
    if re.fullmatch(r"-?\d{1,3}(\.\d{3})*,\d+|-?\d+,\d+", t):
        return float(t.replace(".", "").replace(",", "."))
    return t


def to_xlsx(blocks: list[dict], title: str, theme: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    head_font = Font(bold=True, color=theme["accent"])
    fill = PatternFill("solid", fgColor=theme["fill"])
    tables = [b for b in blocks if b["type"] == "table"]
    used: set[str] = set()
    for n, t in enumerate(tables, 1):
        # a heading right before the table names the sheet
        idx = blocks.index(t)
        name = next((plain(b["runs"]) for b in reversed(blocks[:idx]) if b["type"] == "h"), "") or f"Tabella {n}"
        name = re.sub(r"[\[\]:*?/\\]", " ", name)[:28].strip() or f"Tabella {n}"
        while name in used:
            name = f"{name[:25]} {n}"
        used.add(name)
        ws = wb.create_sheet(name)
        widths: dict[int, int] = {}
        for ri, row in enumerate(t["rows"], 1):
            for ci, c in enumerate(row, 1):
                text = plain(c["runs"])
                cell = ws.cell(row=ri, column=ci, value=_cell_value(text))
                if c["header"]:
                    cell.font, cell.fill = head_font, fill
                widths[ci] = max(widths.get(ci, 8), min(60, len(text) + 2))
        for ci, w in widths.items():
            ws.column_dimensions[get_column_letter(ci)].width = w
        ws.freeze_panes = "A2" if any(c["header"] for c in t["rows"][0]) else None
    rest = [b for b in blocks if b["type"] != "table"]
    if rest or not tables:
        ws = wb.create_sheet("Note", 0 if not tables else None)
        ws.cell(row=1, column=1, value=title).font = Font(bold=True, size=14, color=theme["accent"])
        r = 3
        for line in text_of(rest).splitlines():
            ws.cell(row=r, column=1, value=line)
            r += 1
        ws.column_dimensions["A"].width = 100
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


RENDERERS = {"pdf": to_pdf, "docx": to_docx, "xlsx": to_xlsx}


def extract(fmt: str, data: bytes) -> tuple[int, str]:
    """(pages or sheets, text) of a finished file — opening it again is half of the check."""
    if fmt == "pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        return len(reader.pages), "\n".join(p.extract_text() or "" for p in reader.pages)
    if fmt == "docx":
        from docx import Document
        doc = Document(io.BytesIO(data))
        cells = [c.text for t in doc.tables for row in t.rows for c in row.cells]
        return 1, "\n".join([p.text for p in doc.paragraphs] + cells)
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True)
    text = [str(v) for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for v in row if v is not None]
    return len(wb.worksheets), "\n".join(text)


def _norm(s: str) -> str:
    return re.sub(r"\s+", " ", s).strip().casefold()


def check(fmt: str, data: bytes, title: str, must: list[str]) -> dict:
    if not data:
        return {"ok": False, "pages": 0, "missing": ["file vuoto"]}
    try:
        pages, text = extract(fmt, data)
    except Exception as exc:
        return {"ok": False, "pages": 0, "missing": [f"il file non si riapre: {exc}"]}
    hay = _norm(text)
    wanted = ([title] if fmt != "xlsx" else []) + [m for m in must if str(m).strip()]   # a sheet may hold only tables
    missing = [m for m in wanted if _norm(m) not in hay]
    return {"ok": pages > 0 and not missing, "pages": pages, "missing": missing}


def render(fmt: str, blocks: list[dict], title: str, theme_key: str) -> bytes:
    if fmt not in RENDERERS:
        raise RenderError("formato: pdf, docx o xlsx")
    theme = THEMES.get(theme_key) or THEMES["moderno"]
    return RENDERERS[fmt](blocks, title, theme)
